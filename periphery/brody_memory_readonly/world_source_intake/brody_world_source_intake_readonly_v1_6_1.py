import argparse
import csv
import hashlib
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime

BOUNDARY = {
    "readonly": True,
    "memory_authority": False,
    "memory_decision": False,
    "allowed_to_decide": False,
    "emits_act": False,
    "kernel_binding": False,
    "x108_merge": False,
    "kernel_mutation": False,
    "x108_mutation": False,
    "decision_authority": "KX108_ONLY",
    "scope": "BRODY_WORLD_SOURCE_INTAKE_PERIPHERY_ONLY"
}

SUPPORTED = {".txt", ".md", ".json", ".jsonl", ".csv", ".pdf", ".docx", ".xlsx", ".xls"}

def safe_name(s):
    s = re.sub(r"[^a-zA-Z0-9_\-]+", "_", str(s)).strip("_")
    return s[:80] or "SOURCE"

def sha256_file(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()

def read_text_any(path):
    p = Path(path)
    for enc in ["utf-8-sig", "utf-8", "cp1252", "latin-1"]:
        try:
            return p.read_text(encoding=enc, errors="replace"), f"TEXT_EXTRACTED_{enc}"
        except Exception:
            pass
    return "", "TEXT_EXTRACT_FAILED"

def extract_pdf(path):
    p = Path(path)
    Reader = None
    lib = None

    try:
        from pypdf import PdfReader
        Reader = PdfReader
        lib = "pypdf"
    except Exception:
        try:
            from PyPDF2 import PdfReader
            Reader = PdfReader
            lib = "PyPDF2"
        except Exception:
            return "", "PDF_EXTRACTOR_UNAVAILABLE", None

    try:
        reader = Reader(str(p))
        chunks = []
        for i, page in enumerate(reader.pages):
            try:
                txt = page.extract_text() or ""
            except Exception:
                txt = ""
            if txt.strip():
                chunks.append(f"\n\n--- PAGE {i+1} ---\n{txt.strip()}")
        text = "\n".join(chunks).strip()
        if len(text) < 30:
            return text, "PDF_TEXT_EXTRACTION_EMPTY_OR_WEAK", lib
        return text, "PDF_TEXT_EXTRACTED", lib
    except Exception as e:
        return "", "PDF_EXTRACT_ERROR_" + type(e).__name__, lib

def extract_docx(path):
    p = Path(path)

    try:
        from docx import Document
        doc = Document(str(p))
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if vals:
                    parts.append(" | ".join(vals))
        text = "\n".join(parts).strip()
        return text, "DOCX_TEXT_EXTRACTED", "python-docx"
    except Exception:
        pass

    try:
        with zipfile.ZipFile(p) as z:
            raw = z.read("word/document.xml")
        root = ET.fromstring(raw)
        texts = []
        for el in root.iter():
            if el.text and el.text.strip():
                texts.append(el.text.strip())
        text = " ".join(texts).strip()
        if text:
            return text, "DOCX_TEXT_EXTRACTED_XML_FALLBACK", "zip+xml"
        return "", "DOCX_TEXT_EXTRACTION_EMPTY", "zip+xml"
    except Exception as e:
        return "", "DOCX_EXTRACT_ERROR_" + type(e).__name__, None

def extract_xlsx(path):
    p = Path(path)

    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"\n## SHEET: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append(" | ".join(vals))
        text = "\n".join(lines).strip()
        if text:
            return text, "XLSX_TEXT_EXTRACTED", "openpyxl"
        return "", "XLSX_TEXT_EXTRACTION_EMPTY", "openpyxl"
    except Exception as e:
        return "", "XLSX_EXTRACT_ERROR_" + type(e).__name__, None

def extract_xls(path):
    try:
        import xlrd
        book = xlrd.open_workbook(str(path))
        lines = []
        for sh in book.sheets():
            lines.append(f"\n## SHEET: {sh.name}")
            for r in range(sh.nrows):
                vals = [str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)]
                vals = [v for v in vals if v]
                if vals:
                    lines.append(" | ".join(vals))
        text = "\n".join(lines).strip()
        if text:
            return text, "XLS_TEXT_EXTRACTED", "xlrd"
        return "", "XLS_TEXT_EXTRACTION_EMPTY", "xlrd"
    except Exception:
        return "", "XLS_EXTRACTOR_UNAVAILABLE", None

def extract_csv(path):
    try:
        lines = []
        with Path(path).open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
            reader = csv.reader(f)
            for row in reader:
                vals = [v.strip() for v in row if v and v.strip()]
                if vals:
                    lines.append(" | ".join(vals))
        return "\n".join(lines).strip(), "CSV_TEXT_EXTRACTED", "csv"
    except Exception as e:
        return "", "CSV_EXTRACT_ERROR_" + type(e).__name__, None

def extract_source(path):
    p = Path(path)
    ext = p.suffix.lower()

    if ext in {".txt", ".md", ".json", ".jsonl"}:
        text, status = read_text_any(p)
        return text, status, "plain-text"

    if ext == ".csv":
        return extract_csv(p)

    if ext == ".pdf":
        return extract_pdf(p)

    if ext == ".docx":
        return extract_docx(p)

    if ext == ".xlsx":
        return extract_xlsx(p)

    if ext == ".xls":
        return extract_xls(p)

    return "", "UNSUPPORTED_EXTENSION", None

def write_normalized_md(out_md, src, label, src_hash, text, extraction_status, extractor, original_ext):
    safe_text = (text or "").replace("```", "` ` `")
    usable = len(safe_text.strip()) >= 30

    body = [
        "# BRODY WORLD SOURCE INTAKE READONLY",
        "",
        f"- label: {label}",
        f"- source_original_path: {src}",
        f"- source_original_sha256: {src_hash}",
        f"- original_extension: {original_ext}",
        f"- extraction_status: {extraction_status}",
        f"- extractor: {extractor}",
        f"- content_usable: {str(usable).lower()}",
        "",
        "## Boundary",
        "",
        "- readonly: true",
        "- memory_authority: false",
        "- memory_decision: false",
        "- decision_authority: KX108_ONLY",
        "- kernel_binding: false",
        "- x108_merge: false",
        "- emits_act: false",
        "",
        "## Extracted text",
        "",
        "```text",
        safe_text[:800000],
        "```",
        ""
    ]

    out_md.write_text("\n".join(body), encoding="utf-8")
    return usable

def iter_sources(source):
    source = Path(source)
    if source.is_file():
        return [source]
    return [
        p for p in source.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED
        and "\\node_modules\\" not in str(p)
        and "\\.venv\\" not in str(p)
        and "\\__pycache__\\" not in str(p)
    ]

def build_intake(source, out_root, label):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    source = Path(source)
    out_root = Path(out_root)

    run_dir = out_root / f"BRODY_WORLD_INTAKE_{label}_{ts}"
    normalized = run_dir / "normalized_md"
    normalized.mkdir(parents=True, exist_ok=True)

    records = []
    skipped = []

    for src in iter_sources(source):
        ext = src.suffix.lower()

        if ext not in SUPPORTED:
            skipped.append({"path": str(src), "reason": "UNSUPPORTED_EXTENSION"})
            continue

        src_hash = sha256_file(src)
        text, extraction_status, extractor = extract_source(src)

        out_md = normalized / f"{safe_name(src.stem)}_{src_hash[:12]}.md"
        content_usable = write_normalized_md(
            out_md=out_md,
            src=src,
            label=label,
            src_hash=src_hash,
            text=text,
            extraction_status=extraction_status,
            extractor=extractor,
            original_ext=ext
        )

        records.append({
            "source": "world_intake",
            "path": str(out_md),
            "source_original_path": str(src),
            "source_original_sha256": src_hash,
            "sha256": sha256_file(out_md),
            "extension": ".md",
            "original_extension": ext,
            "snippet": (text or "")[:700],
            "extracted_chars": len(text or ""),
            "content_usable": content_usable,
            "extraction_status": extraction_status,
            "extractor": extractor,
            "world_intake_signal": True,
            "memory_signal": True,
            "zip2_signal": False,
            "graph_signal": False,
            **BOUNDARY
        })

    manifest = {
        "status": "BRODY_WORLD_SOURCE_INTAKE_READONLY_PASS",
        "date": ts,
        "source": str(source),
        "run_dir": str(run_dir),
        "files_count": len(records),
        "skipped_count": len(skipped),
        "usable_files_count": len([r for r in records if r.get("content_usable")]),
        "files": records,
        "skipped": skipped,
        **BOUNDARY
    }

    manifest_path = run_dir / "BRODY_WORLD_SOURCE_INTAKE_MANIFEST.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    current = out_root / "CURRENT_BRODY_WORLD_SOURCE_INTAKE_READONLY.txt"
    current.write_text(
        "\n".join([
            f"CURRENT_BRODY_WORLD_SOURCE_INTAKE_READONLY={run_dir}",
            f"MANIFEST={manifest_path}",
            f"FILES_COUNT={len(records)}",
            f"USABLE_FILES_COUNT={manifest['usable_files_count']}",
            f"SKIPPED_COUNT={len(skipped)}",
            "STATUS=BRODY_WORLD_SOURCE_INTAKE_READONLY_PASS",
            "MEMORY_DECISION=false",
            "DECISION_AUTHORITY=KX108_ONLY"
        ]) + "\n",
        encoding="utf-8"
    )

    print(json.dumps({
        "status": "BRODY_WORLD_SOURCE_INTAKE_READONLY_PASS",
        "run_dir": str(run_dir),
        "manifest": str(manifest_path),
        "files_count": len(records),
        "usable_files_count": manifest["usable_files_count"],
        "skipped_count": len(skipped),
        **BOUNDARY
    }, indent=2, ensure_ascii=False))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True)
    ap.add_argument("--out-root", required=True)
    ap.add_argument("--label", default="SOURCE")
    args = ap.parse_args()

    build_intake(Path(args.source), Path(args.out_root), safe_name(args.label))

if __name__ == "__main__":
    main()
